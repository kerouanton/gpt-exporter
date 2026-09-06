"""Audit semantic preservation from immutable legacy DOCX to normalized DOCX.

This intentionally does NOT compare pagination or Word-specific styling. It
checks only information that GPT Exporter's modern Markdown -> DOCX pipeline
can represent: text structure, headings, lists, emphasis, hyperlinks, tables,
and recoverable embedded assets.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter
from pathlib import Path

from docx import Document


NORMALIZED_SUFFIX = " [normalized]"
IMAGE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
OLE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject"
PACKAGE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/package"


def _sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _style_name(paragraph) -> str:
    try:
        return str(paragraph.style.name or "")
    except (AttributeError, KeyError):
        return ""


def _heading_count(document: Document) -> int:
    return sum(1 for p in document.paragraphs if _style_name(p).casefold().startswith(("heading ", "titre ")))


def _list_count(document: Document) -> int:
    count = 0
    for paragraph in document.paragraphs:
        p_pr = paragraph._p.pPr
        has_num = p_pr is not None and p_pr.numPr is not None
        style = _style_name(paragraph).casefold()
        if has_num or style.startswith(("list bullet", "list number", "liste à puces", "liste num")):
            count += 1
    return count


def _all_paragraphs(document: Document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _hyperlink_count(document: Document) -> int:
    return sum(len(p._p.xpath(".//w:hyperlink")) for p in _all_paragraphs(document))


def _emphasis_phrases(document: Document, attribute: str) -> Counter[str]:
    phrases: Counter[str] = Counter()
    for paragraph in _all_paragraphs(document):
        for run in paragraph.runs:
            if getattr(run, attribute) is True:
                text = re.sub(r"\s+", " ", str(run.text or "")).strip()
                if len(text) >= 2:
                    phrases[text] += 1
    return phrases


def _missing_emphasis(source: Counter[str], normalized: Counter[str]) -> list[str]:
    # Run boundaries are allowed to change through Markdown. Exact multiplicity
    # is therefore advisory; only phrases absent altogether are reported.
    normalized_text = "\n".join(normalized.elements()).casefold()
    missing = [text for text in source if text.casefold() not in normalized_text]
    return missing[:20]


def _literal_br_cells(document: Document) -> int:
    return sum(
        1
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        if "<br>" in cell.text.casefold()
    )


def _relationship_payloads(path: Path) -> tuple[list[str], list[str]]:
    document = Document(path)
    images: list[str] = []
    attachments: list[str] = []
    seen: set[tuple[str, str]] = set()
    for rel in document.part.rels.values():
        reltype = str(rel.reltype)
        if reltype not in {IMAGE_REL, OLE_REL, PACKAGE_REL}:
            continue
        target = getattr(rel, "target_part", None)
        blob = getattr(target, "blob", None)
        if not isinstance(blob, bytes):
            continue
        digest = _sha256_bytes(blob)
        key = (reltype, digest)
        if key in seen:
            continue
        seen.add(key)
        if reltype == IMAGE_REL:
            images.append(digest)
        else:
            attachments.append(digest)
    return images, attachments


def _asset_hashes(asset_dir: Path) -> set[str]:
    hashes: set[str] = set()
    if asset_dir.is_dir():
        for path in asset_dir.rglob("*"):
            if path.is_file():
                try:
                    hashes.add(_sha256_file(path))
                except OSError:
                    pass
    return hashes


def _normalized_path(output_dir: Path, source: Path) -> Path:
    return output_dir / f"{source.stem}{NORMALIZED_SUFFIX}.docx"


def audit_pair(source: Path, normalized: Path, output_dir: Path) -> dict[str, object]:
    source_doc = Document(source)
    normalized_doc = Document(normalized)

    source_images, source_attachments = _relationship_payloads(source)
    normalized_images, _ = _relationship_payloads(normalized)

    source_sha = _sha256_file(source)
    asset_dir = output_dir / "assets" / "legacy" / source_sha[:12]
    exported_hashes = _asset_hashes(asset_dir)

    source_bold = _emphasis_phrases(source_doc, "bold")
    normalized_bold = _emphasis_phrases(normalized_doc, "bold")
    source_italic = _emphasis_phrases(source_doc, "italic")
    normalized_italic = _emphasis_phrases(normalized_doc, "italic")

    source_tables = len(source_doc.tables)
    normalized_tables = len(normalized_doc.tables)
    source_headings = _heading_count(source_doc)
    normalized_headings = _heading_count(normalized_doc)
    source_lists = _list_count(source_doc)
    normalized_lists = _list_count(normalized_doc)
    source_links = _hyperlink_count(source_doc)
    normalized_links = _hyperlink_count(normalized_doc)
    literal_br_cells = _literal_br_cells(normalized_doc)

    missing_asset_hashes = sorted((set(source_images) | set(source_attachments)) - exported_hashes)
    missing_bold = _missing_emphasis(source_bold, normalized_bold)
    missing_italic = _missing_emphasis(source_italic, normalized_italic)

    failures: list[str] = []
    warnings: list[str] = []
    if normalized_tables < source_tables:
        failures.append(f"tables {normalized_tables} < {source_tables}")
    if normalized_headings < source_headings:
        failures.append(f"headings {normalized_headings} < {source_headings}")
    if normalized_lists < source_lists:
        failures.append(f"lists {normalized_lists} < {source_lists}")
    if source_images and len(normalized_images) < len(source_images):
        failures.append(f"embedded images {len(normalized_images)} < {len(source_images)}")
    if missing_asset_hashes:
        failures.append(f"{len(missing_asset_hashes)} source asset payload(s) missing from exported assets")
    if normalized_links < source_links:
        warnings.append(f"hyperlinks {normalized_links} < {source_links}")
    if missing_bold:
        warnings.append(f"bold phrases not found: {len(missing_bold)}")
    if missing_italic:
        warnings.append(f"italic phrases not found: {len(missing_italic)}")
    if literal_br_cells:
        warnings.append(f"literal <br> text in {literal_br_cells} table cell(s)")

    status = "FAIL" if failures else ("WARN" if warnings else "PASS")
    return {
        "source": source.name,
        "normalized": normalized.name,
        "status": status,
        "source_tables": source_tables,
        "normalized_tables": normalized_tables,
        "source_headings": source_headings,
        "normalized_headings": normalized_headings,
        "source_lists": source_lists,
        "normalized_lists": normalized_lists,
        "source_hyperlinks": source_links,
        "normalized_hyperlinks": normalized_links,
        "source_images": len(source_images),
        "normalized_images": len(normalized_images),
        "source_attachments": len(source_attachments),
        "exported_asset_files": len([p for p in asset_dir.rglob("*") if p.is_file()]) if asset_dir.is_dir() else 0,
        "literal_br_cells": literal_br_cells,
        "missing_bold_examples": missing_bold,
        "missing_italic_examples": missing_italic,
        "failures": failures,
        "warnings": warnings,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Audit semantic parity of legacy normalized DOCX files")
    parser.add_argument("--source-dir", type=Path, required=True, help="Directory containing immutable historical DOCX files")
    parser.add_argument("--normalized-dir", type=Path, required=True, help="Directory containing [normalized].docx files and assets/")
    parser.add_argument("--json", type=Path, default=Path("legacy-semantic-audit.json"))
    parser.add_argument("--csv", type=Path, default=Path("legacy-semantic-audit.csv"))
    args = parser.parse_args(argv)

    source_dir = args.source_dir.expanduser().resolve()
    output_dir = args.normalized_dir.expanduser().resolve()
    sources = sorted(path for path in source_dir.glob("*.docx") if NORMALIZED_SUFFIX not in path.stem)

    results: list[dict[str, object]] = []
    missing_normalized: list[str] = []
    for source in sources:
        normalized = _normalized_path(output_dir, source)
        if not normalized.is_file():
            missing_normalized.append(normalized.name)
            continue
        result = audit_pair(source, normalized, output_dir)
        results.append(result)
        print(f"{result['status']:4}  {source.name}")
        for problem in result["failures"]:
            print(f"      FAIL: {problem}")
        for warning in result["warnings"]:
            print(f"      WARN: {warning}")

    summary = {
        "source_count": len(sources),
        "audited_count": len(results),
        "missing_normalized": missing_normalized,
        "pass": sum(result["status"] == "PASS" for result in results),
        "warn": sum(result["status"] == "WARN" for result in results),
        "fail": sum(result["status"] == "FAIL" for result in results),
    }
    payload = {"schema": "gpt-exporter-legacy-semantic-audit-v1", "summary": summary, "results": results}
    args.json.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    fieldnames = [
        "source", "normalized", "status", "source_tables", "normalized_tables",
        "source_headings", "normalized_headings", "source_lists", "normalized_lists",
        "source_hyperlinks", "normalized_hyperlinks", "source_images", "normalized_images",
        "source_attachments", "exported_asset_files", "literal_br_cells", "failures", "warnings",
    ]
    with args.csv.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for result in results:
            row = {name: result.get(name, "") for name in fieldnames}
            row["failures"] = "; ".join(result["failures"])
            row["warnings"] = "; ".join(result["warnings"])
            writer.writerow(row)

    print()
    print("Semantic audit summary")
    print("======================")
    print(f"Sources : {summary['source_count']}")
    print(f"Audited : {summary['audited_count']}")
    print(f"PASS    : {summary['pass']}")
    print(f"WARN    : {summary['warn']}")
    print(f"FAIL    : {summary['fail']}")
    print(f"Missing : {len(missing_normalized)}")
    print(f"JSON    : {args.json.resolve()}")
    print(f"CSV     : {args.csv.resolve()}")

    return 1 if summary["fail"] or missing_normalized else 0


if __name__ == "__main__":
    raise SystemExit(main())
