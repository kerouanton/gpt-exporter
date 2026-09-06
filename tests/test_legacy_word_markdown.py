import os
file_name = os.path.basename(__file__)
print(f"The filename of this script is: {file_name}")

import tempfile
import unittest
from pathlib import Path

from docx import Document

from gpt_exporter.legacy.word_markdown import source_block_markdown


class LegacyWordMarkdownTests(unittest.TestCase):
    def test_preserves_heading_list_emphasis_break_and_table_structure(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "legacy.docx"
            document = Document()
            document.add_heading("Structured heading", level=2)

            bullet = document.add_paragraph(style="List Bullet")
            bullet.add_run("bullet ")
            bullet.add_run("bold").bold = True

            numbered = document.add_paragraph(style="List Number")
            numbered.add_run("numbered ")
            numbered.add_run("italic").italic = True

            paragraph = document.add_paragraph()
            paragraph.add_run("first line")
            paragraph.add_run().add_break()
            paragraph.add_run("second line")

            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Key"
            table.cell(0, 1).text = "Value"
            table.cell(1, 0).text = "A"
            cell = table.cell(1, 1)
            cell.text = ""
            cell.paragraphs[0].add_run("important").bold = True
            cell.paragraphs[0].add_run(" value")
            document.save(source)

            blocks = source_block_markdown(source)
            rendered = "\n\n".join(blocks.values())

            self.assertIn("## Structured heading", rendered)
            self.assertIn("- bullet **bold**", rendered)
            self.assertIn("1. numbered *italic*", rendered)
            self.assertIn("first line  \nsecond line", rendered)
            self.assertIn("| Key | Value |", rendered)
            self.assertIn("| A | **important** value |", rendered)


if __name__ == "__main__":
    unittest.main()
