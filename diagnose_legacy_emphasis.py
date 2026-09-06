"""Inspect Word run segmentation and Markdown-fence hazards in legacy DOCX."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document


def _style_key(style) -> str:
    try:
        return str(style.style_id or style.name or "")
    except (AttributeError, KeyError):
        return ""


def _style_flag(style, attribute: str) -> bool | None:
    visited: set[str] = set()
    current = style
    while current is not None:
        key = _style_key(current)
        if key in visited:
            break
        visited.add(key)
        try:
            value = getattr(current.font, attribute)
        except (AttributeError, KeyError):
            value = None
        if value is not None:
            return bool(value)
        try:
            current = current.base_style
        except (AttributeError, KeyError):
            current = None
    return None


def _effective_inline(run, attribute: str) -> bool:
    direct = getattr(run, attribute)
    if direct is not None:
        return bool(direct)
    try:
        styled = _style_flag(run.style, attribute)
    except (AttributeError, KeyError):
        styled = None
    return bool(styled) if styled is not None else False


def _all_paragraphs(document: Document):
    yield "body", document.paragraphs
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                yield f"table[{table_index}] row[{row_index}] cell[{cell_index}]", cell.paragraphs


def _normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip().casefold()


def _print_paragraph(location: str, paragraph_index: int, paragraph) -> None:
    print(f"{location} paragraph[{paragraph_index}]")
    print(f"PARAGRAPH STYLE: {_style_key(paragraph.style)!r}")
    print(f"TEXT: {paragraph.text!r}")
    for index, run in enumerate(paragraph.runs):
        if not run.text:
            continue
        try:
            style_id = _style_key(run.style)
            style_name = str(run.style.name or "")
        except (AttributeError, KeyError):
            style_id = ""
            style_name = ""
        print(
            f"  run[{index:02}] text={run.text!r} "
            f"direct_bold={run.bold!r} direct_italic={run.italic!r} "
            f"style_id={style_id!r} style_name={style_name!r} "
            f"effective_bold={_effective_inline(run, 'bold')} "
            f"effective_italic={_effective_inline(run, 'italic')}"
        )


def inspect(path: Path, phrase: str) -> int:
    document = Document(path)
    needle = _normalize(phrase)
    found = 0
    print(f"FILE: {path}")
    for location, paragraphs in _all_paragraphs(document):
        for paragraph_index, paragraph in enumerate(paragraphs):
            if needle not in _normalize(paragraph.text):
                continue
            found += 1
            print()
            print(f"MATCH {found}:")
            _print_paragraph(location, paragraph_index, paragraph)
    if not found:
        print("NO MATCH")
    return found


def inspect_fences(path: Path) -> int:
    """Show paragraphs containing literal Markdown fence markers and their Word style."""
    document = Document(path)
    found = 0
    print(f"FILE: {path}")
    for location, paragraphs in _all_paragraphs(document):
        for paragraph_index, paragraph in enumerate(paragraphs):
            if "```" not in paragraph.text and "~~~" not in paragraph.text:
                continue
            found += 1
            print()
            print(f"FENCE {found}:")
            _print_paragraph(location, paragraph_index, paragraph)
    if not found:
        print("NO FENCES")
    return found


def main() -> int:
    parser = argparse.ArgumentParser(description="Inspect legacy Word emphasis or Markdown-fence hazards")
    parser.add_argument("source", type=Path)
    parser.add_argument("phrase", nargs="?")
    parser.add_argument("--normalized", type=Path)
    parser.add_argument("--fences", action="store_true", help="List paragraphs containing literal ``` or ~~~ markers")
    args = parser.parse_args()

    source = args.source.expanduser().resolve()
    normalized = args.normalized.expanduser().resolve() if args.normalized is not None else None

    if args.fences:
        inspect_fences(source)
        if normalized is not None:
            print("\n" + "=" * 80 + "\n")
            inspect_fences(normalized)
        return 0

    if not args.phrase:
        parser.error("phrase is required unless --fences is used")

    inspect(source, args.phrase)
    if normalized is not None:
        print("\n" + "=" * 80 + "\n")
        inspect(normalized, args.phrase)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
